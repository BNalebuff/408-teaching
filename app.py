#!/usr/bin/env python3
"""
Negotiation Journal Feedback - Flask App
Local:  python app.py  →  http://localhost:8765
Deploy: Render.com (see README)
"""

import os, json, io, base64, tempfile, traceback
from flask import Flask, request, jsonify, send_from_directory
import requests

app = Flask(__name__, static_folder='static')
API_KEY = os.environ.get("ANTHROPIC_API_KEY", "")

# ── Routes ────────────────────────────────────────────────

@app.route('/')
def index():
    return send_from_directory('static', 'index.html')

@app.route('/api/feedback', methods=['POST'])
def feedback():
    try:
        data = request.get_json()
        result = call_anthropic(data['payload'])
        return jsonify({'ok': True, 'data': result})
    except Exception as e:
        traceback.print_exc()
        return jsonify({'ok': False, 'error': str(e)}), 500

@app.route('/api/annotate', methods=['POST'])
def annotate():
    """Accept uploaded file + feedback JSON, return annotated file."""
    try:
        file      = request.files.get('file')
        feedback  = json.loads(request.form.get('feedback', '{}'))
        student   = request.form.get('student', 'Student')
        grade     = request.form.get('grade', '')

        if not file:
            return jsonify({'ok': False, 'error': 'No file uploaded'}), 400

        filename = file.filename.lower()
        raw      = file.read()

        if filename.endswith('.docx'):
            out_bytes, out_name = annotate_docx(raw, feedback, student, grade)
            mime = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        else:
            # PDF or anything else → produce annotated PDF
            out_bytes, out_name = annotate_pdf(raw, feedback, student, grade)
            mime = 'application/pdf'

        b64 = base64.b64encode(out_bytes).decode('utf-8')
        return jsonify({'ok': True, 'filename': out_name, 'mime': mime, 'data': b64})

    except Exception as e:
        traceback.print_exc()
        return jsonify({'ok': False, 'error': str(e)}), 500


# ── Anthropic API ─────────────────────────────────────────

def call_anthropic(payload):
    if not API_KEY:
        raise ValueError(
            "ANTHROPIC_API_KEY is not set. "
            "Set it as an environment variable and restart."
        )
    resp = requests.post(
        'https://api.anthropic.com/v1/messages',
        json=payload,
        headers={
            'Content-Type':      'application/json',
            'x-api-key':         API_KEY,
            'anthropic-version': '2023-06-01',
        },
        timeout=60
    )
    if not resp.ok:
        raise ValueError(f"Anthropic API error {resp.status_code}: {resp.text}")
    return resp.json()


# ── DOCX annotation ───────────────────────────────────────

def annotate_docx(raw_bytes, feedback, student, grade):
    """Add a styled feedback section at the end of the DOCX."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import copy

    doc = Document(io.BytesIO(raw_bytes))

    # Divider
    doc.add_paragraph()
    div = doc.add_paragraph('─' * 60)
    div.runs[0].font.color.rgb = RGBColor(0x8b, 0x2e, 0x0f)

    # Header
    hdr = doc.add_paragraph()
    hdr.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = hdr.add_run(f'Instructor Feedback — {student}')
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x1a, 0x16, 0x12)

    if grade:
        g = doc.add_paragraph()
        gr = g.add_run(f'Grade: {grade}/100')
        gr.bold = True
        gr.font.size = Pt(12)
        gr.font.color.rgb = RGBColor(0x8b, 0x2e, 0x0f)

    doc.add_paragraph()

    LABELS = {
        'pie':        '◉ Pie Identification & Calculation',
        'batna':      '⊥ BATNA Awareness',
        'tree':       '⌘ Creative Value / The Tree',
        'nexxtoil':   '⚑ Nexxtoil: Payment',
        'principled': '≡ Principled Reasoning',
        'reflection': '◎ Depth of Reflection',
        'roleplay':   '◈ Playing the Assigned Role',
        'prep':       '▣ Preparation',
        'overall':    '★ Overall Assessment',
    }

    for key, label in LABELS.items():
        text = feedback.get(key)
        if not text:
            continue
        # Section heading
        lp = doc.add_paragraph()
        lr = lp.add_run(label)
        lr.bold = True
        lr.font.size = Pt(11)
        lr.font.color.rgb = RGBColor(0x8b, 0x2e, 0x0f)
        # Body
        bp = doc.add_paragraph(text)
        bp.runs[0].font.size = Pt(11)
        doc.add_paragraph()

    # Strengths / improvements
    strengths    = feedback.get('strengths', [])
    improvements = feedback.get('improvements', [])

    if strengths:
        sp = doc.add_paragraph()
        sr = sp.add_run('Strengths')
        sr.bold = True; sr.font.size = Pt(11)
        sr.font.color.rgb = RGBColor(0x2a, 0x5c, 0x3f)
        for s in strengths:
            p = doc.add_paragraph(f'✓  {s}')
            p.runs[0].font.size = Pt(11)
            p.runs[0].font.color.rgb = RGBColor(0x2a, 0x5c, 0x3f)
            p.paragraph_format.left_indent = Pt(16)

    if improvements:
        doc.add_paragraph()
        ip = doc.add_paragraph()
        ir = ip.add_run('Areas to Develop')
        ir.bold = True; ir.font.size = Pt(11)
        ir.font.color.rgb = RGBColor(0x8b, 0x2e, 0x0f)
        for i in improvements:
            p = doc.add_paragraph(f'→  {i}')
            p.runs[0].font.size = Pt(11)
            p.runs[0].font.color.rgb = RGBColor(0x8b, 0x2e, 0x0f)
            p.paragraph_format.left_indent = Pt(16)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    safe_name = student.replace(' ', '_')
    return buf.read(), f'{safe_name}_feedback.docx'


# ── PDF annotation ────────────────────────────────────────

def annotate_pdf(raw_bytes, feedback, student, grade):
    """Append a feedback page to the PDF using reportlab + pypdf."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    from reportlab.lib.units import inch
    from pypdf import PdfReader, PdfWriter

    ACCENT  = colors.HexColor('#8b2e0f')
    GREEN   = colors.HexColor('#2a5c3f')
    INK     = colors.HexColor('#1a1612')
    MUTED   = colors.HexColor('#6b5f4e')

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('FBTitle',
        fontSize=16, leading=20, textColor=INK,
        fontName='Helvetica-Bold', spaceAfter=6)
    grade_style = ParagraphStyle('FBGrade',
        fontSize=12, leading=16, textColor=ACCENT,
        fontName='Helvetica-Bold', spaceAfter=12)
    section_style = ParagraphStyle('FBSection',
        fontSize=11, leading=14, textColor=ACCENT,
        fontName='Helvetica-Bold', spaceBefore=12, spaceAfter=4)
    body_style = ParagraphStyle('FBBody',
        fontSize=10, leading=14, textColor=INK,
        fontName='Helvetica', spaceAfter=6)
    bullet_pos_style = ParagraphStyle('FBBulletPos',
        fontSize=10, leading=14, textColor=GREEN,
        fontName='Helvetica', leftIndent=16, spaceAfter=3)
    bullet_imp_style = ParagraphStyle('FBBulletImp',
        fontSize=10, leading=14, textColor=ACCENT,
        fontName='Helvetica', leftIndent=16, spaceAfter=3)

    LABELS = {
        'pie':        '&#9711; Pie Identification &amp; Calculation',
        'batna':      'BATNA Awareness',
        'tree':       'Creative Value / The Tree',
        'nexxtoil':   'Nexxtoil: Payment',
        'principled': 'Principled Reasoning',
        'reflection': 'Depth of Reflection',
        'roleplay':   'Playing the Assigned Role',
        'prep':       'Preparation',
        'overall':    '&#9733; Overall Assessment',
    }

    story = []
    story.append(Paragraph(f'Instructor Feedback &mdash; {student}', title_style))
    story.append(HRFlowable(width='100%', thickness=2, color=ACCENT, spaceAfter=8))

    if grade:
        story.append(Paragraph(f'Grade: {grade}/100', grade_style))

    for key, label in LABELS.items():
        text = feedback.get(key)
        if not text:
            continue
        # Escape special chars for reportlab XML
        safe = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
        story.append(Paragraph(label, section_style))
        story.append(Paragraph(safe, body_style))

    strengths    = feedback.get('strengths', [])
    improvements = feedback.get('improvements', [])

    if strengths:
        story.append(Spacer(1, 8))
        story.append(Paragraph('<b>Strengths</b>', ParagraphStyle('SH',
            fontSize=11, textColor=GREEN, fontName='Helvetica-Bold', spaceAfter=4)))
        for s in strengths:
            safe = s.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')
            story.append(Paragraph(f'&#10003;  {safe}', bullet_pos_style))

    if improvements:
        story.append(Spacer(1, 8))
        story.append(Paragraph('<b>Areas to Develop</b>', ParagraphStyle('IH',
            fontSize=11, textColor=ACCENT, fontName='Helvetica-Bold', spaceAfter=4)))
        for i in improvements:
            safe = i.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')
            story.append(Paragraph(f'&#8594;  {safe}', bullet_imp_style))

    # Build feedback page PDF in memory
    fb_buf = io.BytesIO()
    doc = SimpleDocTemplate(fb_buf, pagesize=letter,
        leftMargin=inch, rightMargin=inch,
        topMargin=0.75*inch, bottomMargin=0.75*inch)
    doc.build(story)
    fb_buf.seek(0)

    # Merge original PDF + feedback page
    writer = PdfWriter()
    original = PdfReader(io.BytesIO(raw_bytes))
    for page in original.pages:
        writer.add_page(page)
    fb_reader = PdfReader(fb_buf)
    for page in fb_reader.pages:
        writer.add_page(page)

    out_buf = io.BytesIO()
    writer.write(out_buf)
    out_buf.seek(0)

    safe_name = student.replace(' ', '_')
    return out_buf.read(), f'{safe_name}_feedback.pdf'


# ── Run ───────────────────────────────────────────────────

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 8765))
    print(f"\n{'='*52}")
    print(f"  Negotiation Journal Feedback")
    print(f"{'='*52}")
    if API_KEY:
        print(f"  API key: {API_KEY[:8]}...{API_KEY[-4:]}")
    else:
        print("  WARNING: ANTHROPIC_API_KEY not set")
    print(f"  Open: http://localhost:{port}\n")
    app.run(host='0.0.0.0', port=port, debug=False)
